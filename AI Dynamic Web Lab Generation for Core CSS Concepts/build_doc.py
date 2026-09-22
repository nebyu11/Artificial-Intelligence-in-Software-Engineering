import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Base Document Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)
font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
title_run = title_p.add_run("Nebyu_AI: Dynamic Web Lab Generation for Core CSS Concepts")
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

# Subtitle / Metadata
sub_p = doc.add_paragraph()
sub_p.add_run("Learner Submission Document | Artificial Intelligence in Software Engineering\n").font.italic = True
sub_p.add_run("Repository: ").font.bold = True
sub_p.add_run("https://github.com/nebyu11/Artificial-Intelligence-in-Software-Engineering\n")
sub_p.add_run("Task Folder: ").font.bold = True
sub_p.add_run("AI Dynamic Web Lab Generation for Core CSS Concepts")

doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    return p

def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    return p

def add_code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

# Section 1: Full Prompt Texts
add_heading_1("1. Full Prompt Texts")

add_heading_2("Prompt 1: Initial Box Model Lab Generation")
add_code_block(
    'Act as a frontend web developer. Using your Canvas tool, Generate an interactive website that can be used for understanding the CSS Box Model and its relationship with the display property.\n\n'
    'The page must have:\n'
    '1. Two div elements, \'Box 1\' and \'Box 2\', so I can see how they interact. \'Box 1\' will be the one we control.\n'
    '2. The CSS must use different background colors for the content area, the padding area, and the margin area of \'Box 1\' (e.g., using background-clip: content-box). The border should be a solid line.\n'
    '3. A control panel with:\n'
    '- Sliders to control the padding, margin, border-width, and width of \'Box 1\'.\n'
    '- Labels next to the sliders that show the current pixel value.\n'
    '- A Dropdown (select) to change the display property of \'Box 1\' to: block, inline-block, and inline.\n'
    '4. JavaScript that listens to all sliders and the dropdown, and updates the CSS properties of \'Box 1\' in real-time.'
)

add_heading_2("Prompt 2: Refinement Prompt (Side-Specific Controls & Corner Radius)")
add_code_block(
    'Implement sliders to adjust the margin, padding, and border for each side (top, right, bottom, left) individually, and add a separate slider for the corner radius.'
)

add_heading_2("Prompt 3: Flexbox and Grid Playground Generation")
add_code_block(
    'Act as a frontend web developer. Using your Canvas tool, generate an interactive website that can be used as a playground for CSS Flexbox and Grid.\n\n'
    'The page should have:\n'
    '1. A `div` element acting as the container.\n'
    '2. Several `div` elements inside acting as the items (e.g., 5 items).\n'
    '3. Dropdown menus (selects) that allow me to change the CSS properties of the container.\n'
    '4. I need to be able to change:\n'
    '- display (to switch between block, flex, and grid)\n'
    '- flex-direction (row, column)\n'
    '- justify-content (flex-start, center, space-between, etc.)\n'
    '- align-items (flex-start, center, stretch, etc.)\n'
    '- grid-template-columns (e.g., 1fr 1fr, 1fr 1fr 1fr)\n'
    '5. The JavaScript must update the container\'s CSS in real-time when I change a dropdown.'
)

# Section 2: Execution Screenshots
add_heading_1("2. Execution Screenshots")

img_dir = r"C:\Users\Administrator\Desktop\Artificial-Intelligence-in-Software-Engineering\AI Dynamic Web Lab Generation for Core CSS Concepts"

add_heading_2("Screenshot 1: Initial Box Model Lab Output")
p_img1 = doc.add_paragraph()
p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img1.add_run().add_picture(os.path.join(img_dir, "initial_box_model.png"), width=Inches(6.0))
add_body("Figure 1: Initial Box Model Lab demonstrating real-time sliders for width, margin, padding, border width, display dropdown, and sibling Box 2 interaction.")

add_heading_2("Screenshot 2: Refined Box Model Lab (Side-Specific Controls)")
p_img2 = doc.add_paragraph()
p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img2.add_run().add_picture(os.path.join(img_dir, "refined_box_model.png"), width=Inches(6.0))
add_body("Figure 2: Refined Box Model Visualizer exhibiting granular top/right/bottom/left controls for padding, margin, border width, border radius, and dynamic CSS inspector.")

add_heading_2("Screenshot 3: Flexbox & Grid Layout Playground")
p_img3 = doc.add_paragraph()
p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img3.add_run().add_picture(os.path.join(img_dir, "flexbox_grid_playground.png"), width=Inches(6.0))
add_body("Figure 3: Interactive Flexbox and Grid Playground showing container controls, 5 item elements, real-time visual updates, and generated CSS rules.")

# Section 3: Reflection and Synthesis
add_heading_1("3. Reflection and Synthesis")

add_body(
    "Learning Efficacy: Rapidly generating and interacting with dynamic web development labs offers a profound pedagogical advantage over traditional static documentation and static diagrams. By providing immediate visual feedback, dynamic labs enable learners to intuitively understand how browser rendering engines compute visual box boundaries and layout placement. Specifically, observing how setting display: inline immediately disables custom width declarations and neutralizes vertical margin pushing provides clarity that static text cannot replicate. Learners directly witness how inline elements wrap within line boxes, how inline-block retains element dimensions while flowing inline, and how block elements enforce line breaks and consume full horizontal space. This active, experimental approach significantly compresses the learning curve and fosters deep mental models of CSS layout mechanics."
)

add_body(
    "AI Workflow & Iterative Refinement: Utilizing a sequential, iterative prompting workflow yields substantial architectural and maintainability advantages over attempting to generate complex tools via a single, monolithic prompt. In real-world software engineering, building minimum viable products (MVPs) prior to adding feature layers minimizes risk, prevents cognitive overload, and enables focused quality assurance. By first establishing a working core Box Model visualizer in Prompt 1, and subsequently introducing side-specific margin, padding, border, and border-radius controls in Prompt 2, the development process remains controlled and structured. This modular approach eliminates AI prompt drift, ensures edge cases are caught early, and mirrors standard agile development cycles."
)

# Section 4: GitHub Repository Update
add_heading_1("4. GitHub Repository & Task Folder Structure")
add_body(
    "Repository URL: https://github.com/nebyu11/Artificial-Intelligence-in-Software-Engineering\n"
    "Task Directory: AI Dynamic Web Lab Generation for Core CSS Concepts\n\n"
    "Included Artifacts:\n"
    "• initial_box_model.html — Initial Box Model & Display Interactive Lab\n"
    "• refined_box_model.html — Refined Box Model Visualizer with Side-Specific Controls\n"
    "• flexbox_grid_playground.html — Visualizing Flexbox & Grid Layout Playground\n"
    "• initial_box_model.png — Screenshot of Initial Box Model Lab\n"
    "• refined_box_model.png — Screenshot of Refined Box Model Lab\n"
    "• flexbox_grid_playground.png — Screenshot of Flexbox/Grid Playground\n"
    "• README.md — Comprehensive Task Documentation"
)

output_docx = os.path.join(img_dir, "Nebyu_AI: Dynamic Web Lab Generation for Core CSS Concepts.docx")
doc.save(output_docx)
print("Saved DOCX successfully at:", output_docx)
